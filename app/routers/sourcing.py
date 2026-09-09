"""제품 소싱 에이전트 라우터.

엑셀/CSV 업로드 → AI 추출 → 가격/마진 계산 → 검수 → 승인 → (엑셀 내보내기).
구글시트 연동은 후속(서비스계정 준비 후). docs/sourcing_agent_design.md 참조.
"""
import io
import logging

from fastapi import APIRouter, Depends, File, Request, UploadFile
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session

from app.auth.dependencies import get_current_user
from app.auth.tenant import get_company_id
from app.database import get_db
from app.models import Product, SourcingBatch
from app.models.user import User
from app.sourcing import pipeline, pricing, sheets_sync

log = logging.getLogger(__name__)

router = APIRouter(prefix="/sourcing")
templates = Jinja2Templates(directory="app/templates")


class ParseError(Exception):
    """업로드 파일 파싱 실패 — 사용자에게 친절한 메시지로 노출."""


def _grid_from_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    grid = [[("" if c.value is None else str(c.value)).strip() for c in row] for row in ws.iter_rows()]
    wb.close()
    return grid


def _grid_from_xls(content: bytes) -> list[list[str]]:
    import xlrd
    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    return [[("" if v is None else str(v)).strip() for v in sheet.row_values(r)] for r in range(sheet.nrows)]


def _grid_from_csv(content: bytes) -> list[list[str]]:
    import csv
    text = content.decode("utf-8-sig", errors="replace")
    return [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]


def _parse_grid(content: bytes, filename: str) -> list[list[str]]:
    """엑셀(.xlsx/.xls)/CSV 바이트 → 2차원 그리드(list[list[str]]).

    시트 전체를 AI가 읽으므로 헤더/표 구조를 가정하지 않는다.
    파싱 실패 시 ParseError를 던져 라우트에서 친절한 메시지로 처리.
    """
    if not content:
        raise ParseError("빈 파일입니다.")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if ext == "xlsx":
            grid = _grid_from_xlsx(content)
        elif ext == "xls":
            grid = _grid_from_xls(content)
        elif ext == "csv":
            grid = _grid_from_csv(content)
        else:
            # 확장자 불명 — xlsx → xls → csv 순서로 시도
            for fn in (_grid_from_xlsx, _grid_from_xls, _grid_from_csv):
                try:
                    grid = fn(content)
                    break
                except Exception:
                    continue
            else:
                raise ParseError("지원하지 않는 파일 형식입니다. .xlsx · .xls · .csv 만 업로드하세요.")
    except ParseError:
        raise
    except Exception:
        raise ParseError(
            "파일을 읽을 수 없습니다. 손상되었거나 형식이 맞지 않습니다. "
            "엑셀에서 '다른 이름으로 저장 → .xlsx' 또는 .csv 로 저장한 뒤 다시 업로드해 주세요."
        )

    if not grid:
        raise ParseError("내용이 비어 있습니다.")
    return grid


@router.get("")
def index(request: Request, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    cid = get_company_id(current_user)
    batches = (
        db.query(SourcingBatch)
        .filter(SourcingBatch.company_id == cid)
        .order_by(SourcingBatch.created_at.desc())
        .limit(50)
        .all()
    )
    return templates.TemplateResponse(request, "sourcing/list.html", {
        "user": current_user, "batches": batches, "active_page": "sourcing",
    })


@router.get("/upload")
def upload_form(request: Request, err: str = "", current_user: User = Depends(get_current_user)):
    return templates.TemplateResponse(request, "sourcing/upload.html", {
        "user": current_user, "active_page": "sourcing", "err": err,
    })


@router.post("/upload")
async def upload(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    cid = get_company_id(current_user)
    content = await file.read()
    try:
        grid = _parse_grid(content, file.filename or "upload.xlsx")
    except ParseError as e:
        from urllib.parse import quote
        return RedirectResponse(f"/sourcing/upload?err={quote(str(e))}", status_code=303)

    batch = SourcingBatch(
        company_id=cid,
        source_filename=file.filename,
        file_type=(file.filename or "").rsplit(".", 1)[-1].lower() if file.filename else None,
        status="uploaded",
        created_by=current_user.id,
    )
    db.add(batch)
    db.commit()

    pipeline.run_extraction(db, batch, grid, company_id=cid)
    return RedirectResponse(f"/sourcing/{batch.id}/review", status_code=303)


@router.get("/{batch_id}/review")
def review(batch_id: str, request: Request, db: Session = Depends(get_db),
           current_user: User = Depends(get_current_user)):
    cid = get_company_id(current_user)
    batch = db.query(SourcingBatch).filter(
        SourcingBatch.id == batch_id, SourcingBatch.company_id == cid).first()
    if not batch:
        return RedirectResponse("/sourcing", status_code=303)
    products = db.query(Product).filter(Product.sourcing_batch_id == batch_id).all()
    rows = [{"product": p, "settlement": pipeline.settlement_table(p)} for p in products]
    return templates.TemplateResponse(request, "sourcing/review.html", {
        "user": current_user, "batch": batch, "rows": rows,
        "scenarios": pricing.DEFAULT_COMMISSION_SCENARIOS, "active_page": "sourcing",
        "sheets_configured": sheets_sync.is_configured(),
    })


@router.post("/product/{product_id}/approve")
def approve(product_id: str, db: Session = Depends(get_db),
            current_user: User = Depends(get_current_user)):
    cid = get_company_id(current_user)
    product = db.query(Product).filter(Product.id == product_id, Product.company_id == cid).first()
    if product:
        product.review_status = "approved"
        product.status = "active"
        db.commit()
    return {"ok": bool(product), "id": product_id, "review_status": "approved"}


@router.post("/product/{product_id}/enrich")
def enrich(product_id: str, db: Session = Depends(get_db),
           current_user: User = Depends(get_current_user)):
    """2차 보강(웹서치→규제표현→카피) 실행. 동기 — 웹서치로 수십 초 소요 가능."""
    cid = get_company_id(current_user)
    product = db.query(Product).filter(Product.id == product_id, Product.company_id == cid).first()
    if not product:
        return {"ok": False, "error": "not found"}
    try:
        pipeline.enrich_product(db, product)
        return {
            "ok": True, "id": product_id,
            "compliance": product.compliance or {},
            "generated_copy": product.generated_copy or {},
            "review_status": product.review_status,
        }
    except Exception as e:
        log.exception("enrich failed")
        return {"ok": False, "error": f"{type(e).__name__}: {e}"}


@router.post("/{batch_id}/sync-sheet")
def sync_sheet(batch_id: str, db: Session = Depends(get_db),
               current_user: User = Depends(get_current_user)):
    """검수 결과를 구글시트에 자동 입력."""
    cid = get_company_id(current_user)
    products = db.query(Product).filter(
        Product.sourcing_batch_id == batch_id, Product.company_id == cid).all()
    if not products:
        return {"ok": False, "error": "상품이 없습니다."}
    result = sheets_sync.sync_batch(batch_id, products)
    if result.get("ok"):
        batch = db.query(SourcingBatch).filter(SourcingBatch.id == batch_id).first()
        if batch:
            from datetime import datetime
            batch.sheet_url = result.get("url")
            batch.synced_at = datetime.utcnow()
            batch.status = "synced"
            db.commit()
    return result


@router.get("/{batch_id}/export")
def export_excel(batch_id: str, db: Session = Depends(get_db),
                 current_user: User = Depends(get_current_user)):
    """검수 결과를 엑셀로 내보내기 (구글시트 연동 전 대체)."""
    cid = get_company_id(current_user)
    products = db.query(Product).filter(
        Product.sourcing_batch_id == batch_id, Product.company_id == cid).all()

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "상품마스터"
    headers = ["상품명", "브랜드", "카테고리", "소비자가", "공급가(내부)", "공구가",
               "할인율", "마진율", "배송비", "출고마감", "A/S", "셀러수수료율", "검수상태"]
    ws.append(headers)
    for p in products:
        ws.append([
            p.name, p.brand, p.category, p.consumer_price, p.supplier_price, p.groupbuy_price,
            p.discount_rate, p.margin_rate, p.shipping_cost, p.dispatch_days, p.as_info,
            p.seller_commission_rate, p.review_status,
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="sourcing_{batch_id[:8]}.xlsx"'},
    )
